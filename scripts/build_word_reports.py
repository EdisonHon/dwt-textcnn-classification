"""Build the Chinese and English professional project reports."""

from __future__ import annotations

import csv
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SUMMARY_CSV = ROOT / "outputs" / "sst2" / "summary" / "formal-5seed" / "summary.csv"
RUNS_CSV = ROOT / "outputs" / "sst2" / "summary" / "formal-5seed" / "runs.csv"
REPORT_DIR = ROOT / "reports"
ASSET_DIR = REPORT_DIR / ".report_assets"

METHODS = ("original", "dwt_ca", "dwt_cd", "pca")
METHOD_LABELS = {
    "zh": {"original": "Original", "dwt_ca": "DWT-cA", "dwt_cd": "DWT-cD", "pca": "PCA"},
    "en": {"original": "Original", "dwt_ca": "DWT-cA", "dwt_cd": "DWT-cD", "pca": "PCA"},
}
COLORS = {
    "original": "#1F4E79",
    "dwt_ca": "#5B9BD5",
    "dwt_cd": "#70AD47",
    "pca": "#ED7D31",
}
NAVY = "1F4E79"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
TEXT = "202020"


def load_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


SUMMARY = {row["method"]: row for row in load_rows(SUMMARY_CSV)}
RUNS = load_rows(RUNS_CSV)


def font(size: int, bold: bool = False, chinese: bool = False) -> ImageFont.FreeTypeFont:
    if chinese:
        name = "msyhbd.ttc" if bold else "msyh.ttc"
    else:
        name = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path("C:/Windows/Fonts") / name), size=size)


def hex_rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))


def draw_center(draw: ImageDraw.ImageDraw, xy: tuple[float, float], text: str, fnt, fill) -> None:
    box = draw.textbbox((0, 0), text, font=fnt)
    draw.text((xy[0] - (box[2] - box[0]) / 2, xy[1] - (box[3] - box[1]) / 2), text, font=fnt, fill=fill)


def save_grouped_chart(lang: str, filename: str) -> Path:
    zh = lang == "zh"
    image = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(image)
    title = "分类性能对比" if zh else "Classification Performance"
    draw.text((110, 55), title, font=font(54, True, zh), fill=hex_rgb(TEXT))
    x0, y0, x1, y1 = 155, 160, 1690, 875
    minimum, maximum = 0.76, 0.84
    for tick in [0.76, 0.78, 0.80, 0.82, 0.84]:
        y = y1 - (tick - minimum) / (maximum - minimum) * (y1 - y0)
        draw.line((x0, y, x1, y), fill="#D9E2F3", width=2)
        draw.text((45, y - 18), f"{tick*100:.0f}%", font=font(28, False, zh), fill="#555555")
    draw.line((x0, y0, x0, y1), fill="#777777", width=3)
    draw.line((x0, y1, x1, y1), fill="#777777", width=3)
    group_width = (x1 - x0) / len(METHODS)
    bar_width = 105
    for index, method in enumerate(METHODS):
        center = x0 + group_width * (index + 0.5)
        values = [float(SUMMARY[method]["test_accuracy_mean"]), float(SUMMARY[method]["test_macro_f1_mean"])]
        stds = [float(SUMMARY[method]["test_accuracy_std"]), float(SUMMARY[method]["test_macro_f1_std"])]
        for j, (value, std) in enumerate(zip(values, stds)):
            left = center - bar_width - 10 + j * (bar_width + 20)
            right = left + bar_width
            top = y1 - (value - minimum) / (maximum - minimum) * (y1 - y0)
            color = hex_rgb(COLORS[method]) if j == 0 else tuple(min(255, c + 40) for c in hex_rgb(COLORS[method]))
            draw.rounded_rectangle((left, top, right, y1), radius=8, fill=color)
            err_top = y1 - (value + std - minimum) / (maximum - minimum) * (y1 - y0)
            err_bottom = y1 - (value - std - minimum) / (maximum - minimum) * (y1 - y0)
            mid = (left + right) / 2
            draw.line((mid, err_top, mid, err_bottom), fill="#202020", width=4)
            draw.line((mid - 18, err_top, mid + 18, err_top), fill="#202020", width=4)
            draw.line((mid - 18, err_bottom, mid + 18, err_bottom), fill="#202020", width=4)
        draw_center(draw, (center, 925), METHOD_LABELS[lang][method], font(30, True, zh), hex_rgb(TEXT))
    legend_y = 985
    for idx, (label, shade) in enumerate((("准确率" if zh else "Accuracy", "#1F4E79"), ("Macro-F1", "#6F91B0"))):
        x = 570 + idx * 360
        draw.rectangle((x, legend_y - 18, x + 42, legend_y + 18), fill=shade)
        draw.text((x + 58, legend_y - 22), label, font=font(28, False, zh), fill=hex_rgb(TEXT))
    output = ASSET_DIR / lang / filename
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, dpi=(220, 220))
    return output


def save_bar_chart(lang: str, filename: str, title: str, key: str, divisor: float, unit: str, decimals: int = 1) -> Path:
    zh = lang == "zh"
    image = Image.new("RGB", (1800, 1000), "white")
    draw = ImageDraw.Draw(image)
    draw.text((110, 50), title, font=font(52, True, zh), fill=hex_rgb(TEXT))
    unit_label = f"单位  {unit}" if zh else f"Unit  {unit}"
    unit_box = draw.textbbox((0, 0), unit_label, font=font(25, False, zh))
    draw.text((1690 - (unit_box[2] - unit_box[0]), 78), unit_label, font=font(25, False, zh), fill="#666666")
    values = [float(SUMMARY[m][key]) / divisor for m in METHODS]
    maximum = max(values) * 1.18
    x0, y0, x1, y1 = 160, 160, 1690, 820
    for step in range(6):
        value = maximum * step / 5
        y = y1 - value / maximum * (y1 - y0)
        draw.line((x0, y, x1, y), fill="#E1E7ED", width=2)
        draw.text((45, y - 18), f"{value:.{decimals}f}", font=font(26, False, zh), fill="#666666")
    draw.line((x0, y0, x0, y1), fill="#777777", width=3)
    draw.line((x0, y1, x1, y1), fill="#777777", width=3)
    slot = (x1 - x0) / len(METHODS)
    for idx, (method, value) in enumerate(zip(METHODS, values)):
        center = x0 + slot * (idx + 0.5)
        left, right = center - 105, center + 105
        top = y1 - value / maximum * (y1 - y0)
        draw.rounded_rectangle((left, top, right, y1), radius=12, fill=hex_rgb(COLORS[method]))
        label = f"{value:.{decimals}f}"
        draw_center(draw, (center, top - 35), label, font(27, True, zh), hex_rgb(TEXT))
        draw_center(draw, (center, 870), METHOD_LABELS[lang][method], font(30, True, zh), hex_rgb(TEXT))
    output = ASSET_DIR / lang / filename
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, dpi=(220, 220))
    return output


def save_training_chart(lang: str) -> Path:
    zh = lang == "zh"
    image = Image.new("RGB", (1800, 1000), "white")
    draw = ImageDraw.Draw(image)
    draw.text((110, 50), "训练效率" if zh else "Training Efficiency", font=font(52, True, zh), fill=hex_rgb(TEXT))
    x0, y0, x1, y1 = 160, 160, 1690, 820
    maximum = 10.5
    for step in range(6):
        value = maximum * step / 5
        y = y1 - value / maximum * (y1 - y0)
        draw.line((x0, y, x1, y), fill="#E1E7ED", width=2)
        draw.text((55, y - 18), f"{value:.1f}", font=font(26, False, zh), fill="#666666")
    slot = (x1 - x0) / len(METHODS)
    for idx, method in enumerate(METHODS):
        center = x0 + slot * (idx + 0.5)
        value = float(SUMMARY[method]["training_seconds_mean"])
        std = float(SUMMARY[method]["training_seconds_std"])
        top = y1 - value / maximum * (y1 - y0)
        draw.rounded_rectangle((center - 95, top, center + 95, y1), radius=10, fill=hex_rgb(COLORS[method]))
        err_top = y1 - (value + std) / maximum * (y1 - y0)
        err_bottom = y1 - max(0, value - std) / maximum * (y1 - y0)
        draw.line((center, err_top, center, err_bottom), fill="#202020", width=4)
        draw.line((center - 20, err_top, center + 20, err_top), fill="#202020", width=4)
        draw.line((center - 20, err_bottom, center + 20, err_bottom), fill="#202020", width=4)
        draw_center(draw, (center, top - 35), f"{value:.2f}", font(27, True, zh), hex_rgb(TEXT))
        draw_center(draw, (center, 875), METHOD_LABELS[lang][method], font(30, True, zh), hex_rgb(TEXT))
    draw.text((x0, 925), "秒，均值 ± 样本标准差" if zh else "seconds, mean ± sample SD", font=font(26, False, zh), fill="#666666")
    output = ASSET_DIR / lang / "training.png"
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, dpi=(220, 220))
    return output


def save_pipeline(lang: str) -> Path:
    zh = lang == "zh"
    image = Image.new("RGB", (1800, 920), "white")
    draw = ImageDraw.Draw(image)
    draw.text((100, 45), "实验流程" if zh else "Experimental Pipeline", font=font(52, True, zh), fill=hex_rgb(TEXT))
    labels = (
        ["SST-2 固定划分", "Train-only 词表", "Word2Vec 300d", "四种嵌入输入", "双通道 MCCNN", "五 seed 评估"]
        if zh
        else ["Fixed SST-2 splits", "Train-only vocabulary", "Word2Vec 300d", "Four embedding inputs", "Two-channel MCCNN", "Five-seed evaluation"]
    )
    centers = [(225 + 270 * idx, 350 + (70 if idx % 2 else 0)) for idx in range(6)]
    for idx, (center, label) in enumerate(zip(centers, labels)):
        x, y = center
        color = "#1F4E79" if idx in (0, 2, 4) else "#5B9BD5"
        draw.rounded_rectangle((x - 180, y - 85, x + 180, y + 85), radius=22, fill=color)
        draw_center(draw, (x, y), label, font(27, True, zh), "white")
        if idx < 5:
            nx, ny = centers[idx + 1]
            draw.line((x + 185, y, nx - 205, ny), fill="#44546A", width=8)
            draw.polygon([(nx - 205, ny), (nx - 235, ny - 18), (nx - 235, ny + 18)], fill="#44546A")
    sub = (
        "Original 300d    DWT-cA 150d    DWT-cD 150d    PCA 150d"
        if not zh
        else "Original 300d    DWT-cA 150d    DWT-cD 150d    PCA 150d"
    )
    draw_center(draw, (900, 690), sub, font(32, True, zh), hex_rgb(NAVY))
    note = "同一数据、词表、网络、训练协议与随机种子" if zh else "Same data, vocabulary, network, training protocol, and seeds"
    draw_center(draw, (900, 780), note, font(29, False, zh), hex_rgb(TEXT))
    output = ASSET_DIR / lang / "pipeline.png"
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, dpi=(220, 220))
    return output


def save_architecture(lang: str) -> Path:
    zh = lang == "zh"
    image = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(image)
    draw.text((100, 45), "MCCNN 双通道结构" if zh else "Two-channel MCCNN Architecture", font=font(52, True, zh), fill=hex_rgb(TEXT))
    labels = {
        "input": "Token ID 序列" if zh else "Token ID sequence",
        "pre": "预训练嵌入通道\n可训练" if zh else "Pretrained channel\ntrainable",
        "rnd": "随机嵌入通道\n可训练" if zh else "Random channel\ntrainable",
        "conv": "共享卷积核\n宽度 5 6 7\n每组 100 filters" if zh else "Shared convolutions\nwidths 5 6 7\n100 filters each",
        "merge": "ReLU 后逐通道求和" if zh else "Sum channels after ReLU",
        "pool": "Max-over-time pooling" if zh else "Max-over-time pooling",
        "out": "Dropout 0.5\n全连接二分类" if zh else "Dropout 0.5\nBinary classifier",
    }
    boxes = {
        "input": (100, 410, 360, 570),
        "pre": (470, 230, 800, 410),
        "rnd": (470, 570, 800, 750),
        "conv": (910, 360, 1240, 620),
        "merge": (1340, 260, 1690, 430),
        "pool": (1340, 505, 1690, 660),
        "out": (1340, 760, 1690, 925),
    }
    for key, rect in boxes.items():
        fill = "#1F4E79" if key in {"input", "conv", "out"} else "#5B9BD5"
        draw.rounded_rectangle(rect, radius=22, fill=fill)
        x = (rect[0] + rect[2]) / 2
        y = (rect[1] + rect[3]) / 2
        lines = labels[key].split("\n")
        for offset, line in enumerate(lines):
            draw_center(draw, (x, y + (offset - (len(lines) - 1) / 2) * 44), line, font(27, True, zh), "white")
    arrows = [((360, 490), (470, 320)), ((360, 490), (470, 660)), ((800, 320), (910, 440)), ((800, 660), (910, 540)), ((1240, 490), (1340, 345)), ((1515, 430), (1515, 505)), ((1515, 660), (1515, 760))]
    for start, end in arrows:
        draw.line((*start, *end), fill="#44546A", width=7)
        angle = math.atan2(end[1] - start[1], end[0] - start[0])
        length = 25
        points = [end, (end[0] - length * math.cos(angle - 0.55), end[1] - length * math.sin(angle - 0.55)), (end[0] - length * math.cos(angle + 0.55), end[1] - length * math.sin(angle + 0.55))]
        draw.polygon(points, fill="#44546A")
    output = ASSET_DIR / lang / "architecture.png"
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, dpi=(220, 220))
    return output


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), LIGHT_GRAY)
        borders.append(tag)


def set_run_font(run, name: str, size: float, bold: bool | None = None, color: str = TEXT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (fld_char, instr, separate, text, end):
        run._r.append(element)


def configure_document(doc: Document, lang: str) -> str:
    zh = lang == "zh"
    body_font = "Microsoft YaHei" if zh else "Arial"
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.different_first_page_header_footer = True
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = body_font
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.22
    for style_name, size, before, after in (("Title", 26, 0, 16), ("Heading 1", 17, 18, 8), ("Heading 2", 13, 12, 5), ("Heading 3", 11, 9, 4)):
        style = styles[style_name]
        style.font.name = body_font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption = styles["Caption"]
    caption.font.name = body_font
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string("555555")
    header = section.header.paragraphs[0]
    header.text = "DWT MCCNN 文本分类项目报告" if zh else "DWT MCCNN Text Classification Project Report"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.runs[0], body_font, 8.5, False, "666666")
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    return body_font


def add_body(doc: Document, text: str, body_font: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, body_font, 10.5, True)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest, body_font, 10.5)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, body_font, 10.5)


def add_bullets(doc: Document, items: list[str], body_font: str) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        set_run_font(run, body_font, 10.2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], body_font: str, widths: list[float] | None = None, font_size: float = 8.6) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[idx])
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        set_run_font(run, body_font, font_size, True, "FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            if widths:
                cell.width = Inches(widths[idx])
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, body_font, font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, path: Path, caption: str, body_font: str, width: float = 6.7) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    inline_shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    inline_shape._inline.docPr.set("descr", caption)
    inline_shape._inline.docPr.set("title", caption)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    run = cap.add_run(caption)
    set_run_font(run, body_font, 9, False, "555555")


def add_section_break(doc: Document) -> None:
    doc.add_page_break()


def performance_rows(lang: str) -> list[list[str]]:
    rows = []
    for method in METHODS:
        item = SUMMARY[method]
        rows.append([
            METHOD_LABELS[lang][method],
            str(int(float(item["embedding_dimension"]))),
            f"{float(item['test_accuracy_mean'])*100:.2f}% ± {float(item['test_accuracy_std'])*100:.2f}%",
            f"{float(item['test_macro_f1_mean'])*100:.2f}% ± {float(item['test_macro_f1_std'])*100:.2f}%",
            "—" if method == "original" else f"{float(item['test_accuracy_difference_from_original'])*100:.2f} pp",
        ])
    return rows


def seed_matrix(lang: str) -> list[list[str]]:
    by_method = {method: {int(row["seed"]): float(row["test_accuracy"]) for row in RUNS if row["method"] == method} for method in METHODS}
    return [[METHOD_LABELS[lang][method], *[f"{by_method[method][seed]*100:.2f}%" for seed in (13, 21, 42, 87, 100)]] for method in METHODS]


def build_report(lang: str, output: Path) -> None:
    zh = lang == "zh"
    doc = Document()
    body_font = configure_document(doc, lang)
    title = "DWT 压缩词嵌入在 MCCNN 文本分类中的实验研究" if zh else "Experimental Study of DWT Compressed Word Embeddings for MCCNN Text Classification"
    subtitle = "SST-2 五随机种子复现与效率评估" if zh else "Five-seed SST-2 Reproduction and Efficiency Evaluation"
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    set_run_font(p.add_run(title), body_font, 26, True, "000000")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(65)
    set_run_font(p.add_run(subtitle), body_font, 14, False, "000000")
    for label, value in (("作者" if zh else "Author", "Tszhei Hon"), ("项目" if zh else "Project", "dwt-textcnn-classification"), ("日期" if zh else "Date", "2026-10-11")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(f"{label}  {value}"), body_font, 11.5, label != ("项目" if zh else "Project"))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement = "专业技术报告" if zh else "Professional Technical Report"
    set_run_font(p.add_run(statement), body_font, 11, True, NAVY)

    add_section_break(doc)
    doc.add_heading("摘要" if zh else "Executive Summary", level=1)
    if zh:
        add_body(doc, "本项目独立实现普通多通道卷积神经网络 MCCNN，并研究将 300 维 Google News Word2Vec 词向量压缩至 150 维后，对 SST-2 情感分类性能与计算资源的影响。实验比较 Original、一级离散小波变换近似系数 DWT-cA、细节系数 DWT-cD 和 PCA 四种输入，严格共享数据划分、train-only 词表、网络结构、训练协议及五个随机种子。", body_font)
        add_body(doc, "五随机种子结果表明，Original 的测试准确率最高，为 81.68% ± 0.90%；PCA 在压缩方法中最好，为 81.33% ± 0.83%，仅低 0.35 个百分点；DWT-cD 和 DWT-cA 分别为 81.04% 和 80.65%。三种 150 维方法均将模型参数、嵌入缓存和 checkpoint 约减半，并将 CUDA 峰值分配显存降低 40.12%。", body_font)
        add_body(doc, "结论是：在当前固定协议下，压缩能够以很小的分类性能代价显著降低资源占用，其中 PCA 提供最佳综合平衡；DWT 成功保留了大部分性能，但没有显示出优于 PCA 或 Original 的证据。Original 平均准确率比参考论文报告的 88.1% 低 6.42 个百分点，该差距应结合论文未完整公开的训练与初始化细节理解。", body_font, "结论是：")
        keywords = "关键词：文本分类；多通道卷积神经网络；离散小波变换；词嵌入压缩；主成分分析；SST-2"
    else:
        add_body(doc, "This project independently implements the ordinary multichannel convolutional neural network MCCNN and evaluates how compressing 300-dimensional Google News Word2Vec vectors to 150 dimensions affects SST-2 sentiment classification and resource use. Original, one-level DWT approximation coefficients, DWT detail coefficients, and PCA are compared under identical data splits, a train-only vocabulary, one network architecture, one training protocol, and five fixed random seeds.", body_font)
        add_body(doc, "Across five seeds, Original achieves the highest test accuracy at 81.68% ± 0.90%. PCA is the strongest compressed representation at 81.33% ± 0.83%, only 0.35 percentage points below Original. DWT-cD and DWT-cA reach 81.04% and 80.65%, respectively. Every 150-dimensional method approximately halves the model parameters, embedding cache, and checkpoint size while reducing peak allocated CUDA memory by 40.12%.", body_font)
        add_body(doc, "The main conclusion is that compression yields substantial efficiency gains for a small accuracy cost, and PCA provides the best overall balance under this fixed protocol. DWT retains most classification performance but does not outperform PCA or Original. The Original mean is 6.42 percentage points below the 88.1% MCCNN result reported by the reference paper, a gap that should be interpreted in light of training and initialization details that were not fully specified.", body_font, "The main conclusion is that ")
        keywords = "Keywords  text classification  multichannel CNN  discrete wavelet transform  embedding compression  PCA  SST-2"
    p = doc.add_paragraph()
    set_run_font(p.add_run(keywords), body_font, 9.5, True, NAVY)

    doc.add_heading("目录" if zh else "Contents", level=1)
    contents = (["1 研究背景与目标", "2 数据与实验设计", "3 方法与实现", "4 正式实验结果", "5 时间 内存与存储", "6 讨论与限制", "7 结论与建议", "参考文献", "附录 五随机种子结果"] if zh else ["1 Background and Objectives", "2 Data and Experimental Design", "3 Methods and Implementation", "4 Formal Experimental Results", "5 Time Memory and Storage", "6 Discussion and Limitations", "7 Conclusions and Recommendations", "References", "Appendix Five-seed Results"])
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(item), body_font, 10.5)

    add_section_break(doc)
    doc.add_heading("1 研究背景与目标" if zh else "1 Background and Objectives", level=1)
    if zh:
        add_body(doc, "卷积神经网络通过局部窗口提取 n-gram 特征，是传统神经文本分类的重要基线。Xiao 等人在 IJCAI 2018 提出的研究以普通 MCCNN 为基线，并进一步研究可变形卷积与池化。其 MCCNN 在 SST-2 上报告 88.1% 准确率。本项目只复现普通 MCCNN，不实现 TF-MCCNN 或 TF-DCNN，以便先建立清晰、可复现的压缩比较基线。", body_font)
        add_body(doc, "词嵌入的维度直接影响嵌入矩阵、卷积核参数、存储和推理开销。Salama 等人在 ACL Findings 2024 研究了使用离散小波变换压缩词与句子嵌入，指出一级 DWT 可将向量分解为低频近似系数 cA 与高频细节系数 cD。受该方法启发，本项目将每个 300 维 Word2Vec 向量视为一维信号，并分别保留 150 维 cA 或 cD，同时使用 PCA 作为同维度对照。", body_font)
        doc.add_heading("1.1 研究问题", level=2)
        add_bullets(doc, ["独立实现的普通 MCCNN 在固定 SST-2 设置下达到什么水平。", "将 Word2Vec 从 300 维压缩到 150 维后，分类性能损失多大。", "DWT-cA、DWT-cD 与 PCA 在相同维度下有何差异。", "压缩能否降低参数量、存储、训练时间、推理延迟和显存。"], body_font)
    else:
        add_body(doc, "Convolutional neural networks extract local n-gram features and remain an important baseline for neural text classification. Xiao et al. introduced transformable convolution and pooling at IJCAI 2018 using an ordinary MCCNN as a baseline. That paper reports 88.1% accuracy on SST-2 for MCCNN. This project reproduces only the ordinary MCCNN, not TF-MCCNN or TF-DCNN, to establish a transparent and reproducible baseline for embedding compression.", body_font)
        add_body(doc, "Embedding dimensionality directly affects lookup-table size, convolution parameters, storage, and inference cost. Salama et al. studied discrete wavelet compression of word and sentence embeddings at Findings of ACL 2024. A one-level DWT separates a vector into low-frequency approximation coefficients cA and high-frequency detail coefficients cD. Following that idea, this project treats each 300-dimensional Word2Vec vector as a one-dimensional signal, retains either 150-dimensional cA or cD, and uses PCA as an equal-dimensional control.", body_font)
        doc.add_heading("1.1 Research Questions", level=2)
        add_bullets(doc, ["What performance does an independent ordinary MCCNN implementation achieve under a fixed SST-2 protocol?", "How much classification performance is lost when Word2Vec is compressed from 300 to 150 dimensions?", "How do DWT-cA, DWT-cD, and PCA differ at the same output dimension?", "Does compression reduce parameter count, storage, training time, inference latency, and GPU memory?"], body_font)
    add_figure(doc, save_pipeline(lang), ("图 1  项目实验流程" if zh else "Figure 1  Project experimental pipeline"), body_font)

    add_section_break(doc)
    doc.add_heading("2 数据与实验设计" if zh else "2 Data and Experimental Design", level=1)
    if zh:
        add_body(doc, "数据来自 Stanford NLP 发布的 2013 年原始 trainDevTestTrees PTB 压缩包。每棵树只使用根节点标签与完整叶子序列，标签 0 和 1 映射为负面，3 和 4 映射为正面，中性标签 2 被移除。最终固定划分为 Train 6,920、Dev 872、Test 1,821，与参考 CNN 论文一致。", body_font)
        add_body(doc, "词表只从 Train 构建，包含 16,282 个普通 token，加上 <pad> 和 <unk> 后共 16,284 项。Dev 与 Test 中未在 Train 出现的 token 映射为 <unk>，避免使用评估数据拟合词表或压缩变换。Word2Vec 精确匹配 14,505 个普通 token，未匹配 token 使用由 seed 控制的 Uniform(-0.25, 0.25) 随机向量。", body_font)
    else:
        add_body(doc, "The dataset is Stanford NLP's original 2013 trainDevTestTrees PTB release. Only the root label and complete leaf sequence of each tree are used. Root labels 0 and 1 map to negative, labels 3 and 4 map to positive, and neutral label 2 is removed. The fixed binary splits contain 6,920 Train, 872 Dev, and 1,821 Test examples, matching the CNN reference paper.", body_font)
        add_body(doc, "The vocabulary is built exclusively from Train. It contains 16,282 lexical tokens plus <pad> and <unk>, for 16,284 entries. Dev and Test tokens absent from Train map to <unk>, preventing evaluation text from fitting the vocabulary or compression transform. Word2Vec exactly matches 14,505 lexical tokens; unmatched tokens use seed-controlled Uniform(-0.25, 0.25) vectors.", body_font)
    add_table(doc, ["项目" if zh else "Item", "固定设置" if zh else "Fixed setting"], [["数据源" if zh else "Source", "Stanford 2013 SST PTB"], ["划分" if zh else "Splits", "6,920 / 872 / 1,821"], ["词表" if zh else "Vocabulary", "Train-only, 16,284"], ["Word2Vec", "Google News 300d"], ["随机种子" if zh else "Seeds", "13, 21, 42, 87, 100"], ["开发集用途" if zh else "Dev usage", "Checkpoint selection only"], ["测试集用途" if zh else "Test usage", "Final evaluation only"]], body_font, [2.0, 4.6], 9.2)
    doc.add_heading("2.1 公平性约束" if zh else "2.1 Fairness Constraints", level=2)
    add_bullets(doc, (["四种方法共享完全相同的数据划分和 token ID。", "每个 seed 的压缩缓存来自该 seed 的同一 Original 矩阵。", "网络结构与训练超参数一致，只有嵌入维度变化。", "开发集选择 checkpoint，测试集不参与调参。", "结论基于五 seed 均值和样本标准差，而非最佳 seed。"] if zh else ["All methods share identical splits and token IDs.", "For each seed, every compressed cache derives from the same Original matrix.", "Network and training hyperparameters are fixed; only embedding dimension changes.", "Dev selects the checkpoint and Test is excluded from tuning.", "Conclusions use five-seed means and sample standard deviations, not the best seed."]), body_font)

    add_section_break(doc)
    doc.add_heading("3 方法与实现" if zh else "3 Methods and Implementation", level=1)
    doc.add_heading("3.1 嵌入表示" if zh else "3.1 Embedding Representations", level=2)
    descriptions = ([
        ("Original", "精确匹配的 Google News Word2Vec 保持 300 维；OOV 与 <unk> 使用 seed 控制的随机初始化。"),
        ("DWT-cA", "使用 PyWavelets、coif1、一级变换和 periodization 边界模式，仅保留 150 维近似系数。"),
        ("DWT-cD", "使用相同 DWT 设置，仅保留 150 维细节系数。"),
        ("PCA", "使用 scikit-learn 完整 SVD，不 whitening，只在 Train 中出现且成功匹配 Word2Vec 的 14,505 个 token 上拟合。"),
    ] if zh else [
        ("Original", "Exact Google News Word2Vec matches remain at 300 dimensions; OOV tokens and <unk> use seed-controlled random initialization."),
        ("DWT-cA", "PyWavelets applies coif1 at one level with periodization boundaries and retains 150 approximation coefficients."),
        ("DWT-cD", "The same DWT configuration retains 150 detail coefficients."),
        ("PCA", "scikit-learn full SVD without whitening is fitted only on the 14,505 Train tokens matched in Word2Vec."),
    ])
    add_table(doc, ["方法" if zh else "Method", "维度" if zh else "Dim.", "说明" if zh else "Definition"], [[name, "300" if name == "Original" else "150", desc] for name, desc in descriptions], body_font, [1.1, 0.7, 4.8], 8.7)
    doc.add_heading("3.2 MCCNN 结构", level=2) if zh else doc.add_heading("3.2 MCCNN Architecture", level=2)
    if zh:
        add_body(doc, "模型包含两个相同维度且均可训练的嵌入通道：预训练通道由对应缓存初始化，随机通道直接在目标维度初始化。三组共享卷积核分别使用宽度 5、6、7，每组 100 个 filters。两个通道分别经过卷积、bias 和 ReLU 后求和，再执行 max-over-time pooling；三个池化向量拼接后通过 dropout 0.5 和全连接层完成二分类。", body_font)
    else:
        add_body(doc, "The model contains two trainable lookup channels with the same dimensionality. The pretrained channel is initialized from the selected cache, while the random channel is initialized directly at the target dimension. Shared convolutions use widths 5, 6, and 7 with 100 filters each. Each channel independently applies convolution, bias, and ReLU before the feature maps are summed. Max-over-time pooling, concatenation, dropout 0.5, and a fully connected layer produce the binary prediction.", body_font)
    add_figure(doc, save_architecture(lang), ("图 2  MCCNN 双通道网络结构" if zh else "Figure 2  Two-channel MCCNN architecture"), body_font)
    doc.add_heading("3.3 训练协议" if zh else "3.3 Training Protocol", level=2)
    add_table(doc, ["参数" if zh else "Parameter", "设置" if zh else "Setting"], [["Batch size", "50"], ["最大 epochs" if zh else "Maximum epochs", "30"], ["优化器" if zh else "Optimizer", "Adadelta"], ["Learning rate", "1.0"], ["rho / eps", "0.95 / 1e-6"], ["Early stopping", "Dev Accuracy, patience 5"], ["初始化" if zh else "Initialization", "Truncated normal, std 0.1"], ["确定性" if zh else "Determinism", "PyTorch deterministic algorithms"]], body_font, [2.2, 4.4], 9.0)

    add_section_break(doc)
    doc.add_heading("4 正式实验结果" if zh else "4 Formal Experimental Results", level=1)
    if zh:
        add_body(doc, "表 4 汇总四种方法在五个固定 seed 上的 Test Accuracy 和 Macro-F1。Original 取得最高平均准确率 81.68%。PCA 仅低 0.35 个百分点，是性能保留最好的压缩方案。DWT-cD 比 DWT-cA 高 0.39 个百分点，但两种 DWT 均未超过 PCA。", body_font)
    else:
        add_body(doc, "Table 4 summarizes Test Accuracy and Macro-F1 across the five fixed seeds. Original has the highest mean accuracy at 81.68%. PCA is only 0.35 percentage points lower and retains the most performance among compressed methods. DWT-cD exceeds DWT-cA by 0.39 percentage points, but neither DWT variant exceeds PCA.", body_font)
    add_table(doc, ["方法" if zh else "Method", "维度" if zh else "Dim.", "Test Accuracy", "Macro-F1", "相对 Original" if zh else "vs. Original"], performance_rows(lang), body_font, [1.05, 0.65, 1.75, 1.75, 1.25], 8.3)
    add_figure(doc, save_grouped_chart(lang, "classification.png"), ("图 3  五 seed 分类性能 均值与样本标准差" if zh else "Figure 3  Five-seed classification performance with sample standard deviation"), body_font)
    if zh:
        add_body(doc, "Original 的 81.68% 平均准确率低于参考论文的 88.1%，差值为 -6.42 个百分点。该结果不是通过挑选最佳 seed 得到，而是完整五 seed 均值。由于参考论文没有完整公开学习率、训练轮数、停止规则和部分初始化细节，本项目将差距作为复现限制报告，而不在观察测试结果后进行针对性调参。", body_font)
    else:
        add_body(doc, "The Original mean of 81.68% is 6.42 percentage points below the reference paper's 88.1%. This value is the complete five-seed mean rather than a selected best run. Because the reference paper does not fully specify learning rate, training duration, stopping rules, and some initialization details, this project reports the gap as a reproduction limitation instead of tuning after observing Test results.", body_font)

    add_section_break(doc)
    doc.add_heading("5 时间 内存与存储" if zh else "5 Time Memory and Storage", level=1)
    doc.add_heading("5.1 参数与文件规模" if zh else "5.1 Parameters and File Size", level=2)
    if zh:
        add_body(doc, "将嵌入从 300 维降至 150 维不仅使两个嵌入通道减半，也使 Conv1d 的输入通道数减半。三种压缩方法的总参数量为 5,156,102，相比 Original 的 10,311,302 减少 50.00%。嵌入缓存约减半至 9.77 MB，checkpoint 约减半至 20.63 MB。", body_font)
    else:
        add_body(doc, "Reducing embeddings from 300 to 150 dimensions halves both lookup channels and halves the Conv1d input dimensionality. Each compressed model has 5,156,102 parameters versus 10,311,302 for Original, a 50.00% reduction. The embedding cache falls to 9.77 MB and the checkpoint to 20.63 MB, both approximately half their Original size.", body_font)
    size_rows = []
    for method in METHODS:
        item = SUMMARY[method]
        size_rows.append([METHOD_LABELS[lang][method], f"{int(float(item['pretrained_embedding_parameters'])):,}", f"{int(float(item['random_embedding_parameters'])):,}", f"{int(float(item['convolution_parameters'])):,}", f"{int(float(item['total_parameters'])):,}", f"{float(item['embedding_cache_bytes'])/1_000_000:.2f} MB", f"{float(item['checkpoint_bytes_mean'])/1_000_000:.2f} MB"])
    add_table(doc, ["方法" if zh else "Method", "预训练嵌入" if zh else "Pretrained", "随机嵌入" if zh else "Random", "卷积" if zh else "Conv.", "总参数" if zh else "Total", "缓存" if zh else "Cache", "Checkpoint"], size_rows, body_font, [0.8, 1.0, 1.0, 0.8, 1.0, 1.0, 1.1], 7.4)
    add_figure(doc, save_bar_chart(lang, "parameters.png", "模型参数量" if zh else "Model Parameter Count", "total_parameters", 1_000_000, "百万参数" if zh else "million parameters", 1), ("图 4  模型总参数量" if zh else "Figure 4  Total model parameter count"), body_font)
    doc.add_heading("5.2 训练效率" if zh else "5.2 Training Efficiency", level=2)
    if zh:
        add_body(doc, "平均总训练时间从 Original 的 8.76 秒降至 DWT-cA 的 7.22 秒、DWT-cD 的 7.11 秒和 PCA 的 6.38 秒。总训练时间受到 early stopping 完成 epoch 数影响，因此平均每 epoch 时间是更直接的计算成本指标：Original 为 0.778 秒，PCA 为 0.643 秒。", body_font)
    else:
        add_body(doc, "Mean total training time falls from 8.76 seconds for Original to 7.22 seconds for DWT-cA, 7.11 seconds for DWT-cD, and 6.38 seconds for PCA. Total duration depends on the number of epochs completed under early stopping, so mean epoch time is the cleaner compute measure: 0.778 seconds for Original and 0.643 seconds for PCA.", body_font)
    add_figure(doc, save_training_chart(lang), ("图 5  总训练时间 均值与样本标准差" if zh else "Figure 5  Total training time with sample standard deviation"), body_font)
    doc.add_page_break()
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)
    doc.add_heading("5.3 推理与内存" if zh else "5.3 Inference and Memory", level=2)
    if zh:
        add_body(doc, "推理使用测试集首个 50 样本 batch，执行 20 次预热和 100 次测量，CUDA 在计时前后同步。Original 平均 batch 延迟为 1.247 ms。压缩方法为 0.933–1.145 ms，但 DWT-cD 和 PCA 的跨 seed 波动较明显，因此结果支持 150 维总体更高效，不支持对三种压缩变换做强速度排名。", body_font)
        add_body(doc, "Original 的 CUDA 峰值分配为 60,389,888 bytes，压缩方法均为 36,159,488 bytes，减少 40.12%。CPU 进程 RSS 还包含 Python、PyTorch、CUDA 库和内存映射，因此未随模型缩小，不应被解释为纯模型内存。", body_font)
    else:
        add_body(doc, "Inference uses the first 50-example Test batch, 20 warm-up iterations, and 100 measured iterations with CUDA synchronization around timing. Original averages 1.247 ms per batch. Compressed methods range from 0.933 to 1.145 ms, but DWT-cD and PCA vary noticeably across seeds. The results support a general efficiency advantage for 150 dimensions, not a strong speed ranking among the three transforms.", body_font)
        add_body(doc, "Peak CUDA allocation is 60,389,888 bytes for Original and 36,159,488 bytes for every compressed method, a 40.12% reduction. Process RSS also includes Python, PyTorch, CUDA libraries, and memory-mapped artifacts. It does not fall with model size and must not be interpreted as model-only CPU memory.", body_font)
    runtime_rows = [[METHOD_LABELS[lang][m], f"{float(SUMMARY[m]['epoch_seconds_mean_mean']):.3f} s", f"{float(SUMMARY[m]['latency_mean_ms_mean']):.3f} ms", f"{float(SUMMARY[m]['throughput_samples_per_second_mean']):,.0f}/s", f"{float(SUMMARY[m]['cuda_peak_allocated_bytes_mean'])/1_048_576:.2f} MiB"] for m in METHODS]
    add_table(doc, ["方法" if zh else "Method", "每 epoch" if zh else "Per epoch", "Batch 延迟" if zh else "Batch latency", "吞吐量" if zh else "Throughput", "CUDA 峰值" if zh else "CUDA peak"], runtime_rows, body_font, [1.1, 1.25, 1.4, 1.45, 1.4], 8.5)
    add_figure(doc, save_bar_chart(lang, "cuda_memory.png", "CUDA 峰值分配显存" if zh else "Peak Allocated CUDA Memory", "cuda_peak_allocated_bytes_mean", 1_048_576, "MiB", 1), ("图 6  CUDA 峰值分配显存" if zh else "Figure 6  Peak allocated CUDA memory"), body_font, width=6.0)

    add_section_break(doc)
    doc.add_heading("6 讨论与限制" if zh else "6 Discussion and Limitations", level=1)
    if zh:
        doc.add_heading("6.1 性能与效率权衡", level=2)
        add_body(doc, "PCA 的平均准确率只比 Original 低 0.35 个百分点，同时获得约 50% 的参数与存储缩减，是本实验最稳健的综合选择。DWT-cD 比 DWT-cA 更好，说明高频变化中确实包含对当前分类任务有用的信息；但 DWT-cD 的跨 seed 标准差为 1.46%，高于 PCA 的 0.83%，稳定性较弱。", body_font)
        add_body(doc, "DWT 的主要优势是无需标签、无需拟合且可以逐向量独立处理。PCA 需要在训练 token 上拟合，但其投影针对当前词向量分布优化。当前结果说明这种数据自适应投影在 SST-2 MCCNN 上更有效。", body_font)
        doc.add_heading("6.2 复现差距", level=2)
        add_body(doc, "参考论文的 88.1% 高于本项目 Original 均值。数据规模与词表数量已与论文对齐，但论文没有公开足以唯一重建训练过程的全部信息。本项目固定 Adadelta 学习率、初始化标准差、early stopping 和 train-only 词表，并完整记录这些假设。因此结果具备内部公平性与可追溯性，但不是对原论文实现的逐比特复刻。", body_font)
        doc.add_heading("6.3 研究限制", level=2)
        add_bullets(doc, ["只使用 SST-2，不能直接外推到其他数据集或任务。", "只使用一级 coif1，没有搜索母小波或分解层级。", "两个嵌入通道训练后会偏离初始压缩表示。", "仅五个 seed，未进行更大规模显著性检验。", "短 GPU benchmark 易受操作系统调度与功耗状态影响。", "CPU RSS 是进程级采样值，不是隔离后的模型内存。"], body_font)
    else:
        doc.add_heading("6.1 Performance Efficiency Tradeoff", level=2)
        add_body(doc, "PCA is only 0.35 percentage points below Original while delivering approximately 50% parameter and storage reductions, making it the strongest overall choice in this experiment. DWT-cD exceeds DWT-cA, indicating that high-frequency variation contains useful information for this classification task. However, DWT-cD has a larger cross-seed standard deviation of 1.46%, compared with 0.83% for PCA.", body_font)
        add_body(doc, "DWT requires no labels, no fitting, and can process vectors independently. PCA must be fitted on training tokens, but its projection adapts to the observed embedding distribution. The present results indicate that this data-adaptive projection is more effective for SST-2 MCCNN.", body_font)
        doc.add_heading("6.2 Reproduction Gap", level=2)
        add_body(doc, "The reference paper's 88.1% exceeds the Original mean in this project. Dataset size and vocabulary scale are aligned, but the paper does not publish enough information to uniquely reconstruct every training decision. This project fixes Adadelta settings, initialization scale, early stopping, and a train-only vocabulary and records those assumptions. The experiment is internally fair and traceable, but it is not a bit-for-bit reconstruction of the authors' implementation.", body_font)
        doc.add_heading("6.3 Limitations", level=2)
        add_bullets(doc, ["Only SST-2 is evaluated, so conclusions may not transfer to other tasks.", "Only one-level coif1 is used; mother wavelet and level are not searched.", "Trainable embedding channels can move away from their initial compressed representation.", "Five seeds do not support broad significance claims.", "Short GPU benchmarks are sensitive to operating-system scheduling and power state.", "Process RSS is sampled at process level and is not isolated model memory."], body_font)

    doc.add_heading("7 结论与建议" if zh else "7 Conclusions and Recommendations", level=1)
    if zh:
        add_body(doc, "本项目完成了从数据、词向量、压缩、MCCNN、五 seed 训练到统一效率测量的完整可复现实验链路。Original 的平均测试准确率最高；PCA 在准确率仅下降 0.35 个百分点的情况下，将模型与存储规模约减半，是当前最佳综合方案。DWT-cD 优于 DWT-cA，但没有超过 PCA。", body_font)
        add_body(doc, "后续工作应与当前正式结果分离并预先定义。优先建议开展训练协议敏感性实验以分析与论文 88.1% 的差距，其次测试不同母小波与分解层级。只有在基础结果稳定后，再考虑 TF-MCCNN、可训练小波或更多数据集。", body_font)
    else:
        add_body(doc, "The project completes a reproducible path from data and pretrained vectors through compression, MCCNN training, five-seed evaluation, and standardized efficiency measurement. Original has the highest mean accuracy. PCA approximately halves model and storage size for a 0.35 percentage-point accuracy cost and is the best overall compressed option. DWT-cD performs better than DWT-cA but does not exceed PCA.", body_font)
        add_body(doc, "Future work should be separated from the formal results and specified in advance. The first priority is a training-protocol sensitivity study to investigate the gap to the paper's 88.1%. The next priority is evaluation of alternative mother wavelets and decomposition levels. TF-MCCNN, trainable wavelets, and additional datasets should be considered only after the baseline is stable.", body_font)

    add_section_break(doc)
    doc.add_heading("参考文献" if zh else "References", level=1)
    references = [
        "Xiao, L., Zhang, H., Chen, W., Wang, Y., and Jin, Y. 2018. Transformable Convolutional Neural Network for Text Classification. Proceedings of IJCAI 2018, 4496–4503.",
        "Salama, R. A., Youssef, A., and Diab, M. 2024. Semantic Word and Sentence Embeddings Compression using Discrete Wavelet Transform. Findings of ACL 2024, 15963–15977.",
        "Socher, R., Perelygin, A., Wu, J., Chuang, J., Manning, C. D., Ng, A., and Potts, C. 2013. Recursive Deep Models for Semantic Compositionality Over a Sentiment Treebank. EMNLP 2013.",
        "Mikolov, T., Chen, K., Corrado, G., and Dean, J. 2013. Efficient Estimation of Word Representations in Vector Space. ICLR Workshop.",
        "Pedregosa, F. et al. 2011. Scikit-learn Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
        "Lee, G. et al. PyWavelets A Python Package for Wavelet Analysis. Journal of Open Source Software.",
    ]
    for idx, reference in enumerate(references, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(f"[{idx}] {reference}"), body_font, 9.5)

    doc.add_heading("附录 五随机种子结果" if zh else "Appendix Five-seed Results", level=1)
    if zh:
        add_body(doc, "下表列出每种方法在五个正式 seed 上的测试准确率。均值与标准差以完整五次运行计算。", body_font)
    else:
        add_body(doc, "The table lists Test Accuracy for every method and formal seed. Reported means and standard deviations use all five runs.", body_font)
    add_table(doc, ["方法" if zh else "Method", "Seed 13", "Seed 21", "Seed 42", "Seed 87", "Seed 100"], seed_matrix(lang), body_font, [1.15, 1.05, 1.05, 1.05, 1.05, 1.1], 8.5)
    doc.add_heading("复现入口" if zh else "Reproduction Entry Points", level=2)
    commands = [
        "dwt-textcnn prepare-data",
        "dwt-textcnn prepare-embeddings --config configs/<method>.yaml",
        "dwt-textcnn train --config configs/<method>.yaml --seed <seed>",
        "dwt-textcnn evaluate --run-dir <run-dir>",
        "dwt-textcnn benchmark --run-dir <run-dir>",
        "dwt-textcnn summarize --config configs/base.yaml --run-id formal-5seed",
    ]
    for command in commands:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(command), "Consolas", 8.8, False, "333333")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.author = "Tszhei Hon"
    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.core_properties.keywords = "MCCNN, DWT, PCA, Word2Vec, SST-2, embedding compression"
    doc.save(output)


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    build_report("zh", REPORT_DIR / "DWT_MCCNN_项目技术报告_中文.docx")
    build_report("en", REPORT_DIR / "DWT_MCCNN_Project_Technical_Report_English.docx")


if __name__ == "__main__":
    main()
